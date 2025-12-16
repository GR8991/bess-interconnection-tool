import pandas as pd
import streamlit as st
from io import BytesIO
from reportlab.pdfgen import canvas
from reportlab.lib.pagesizes import letter
from docx import Document


# =====================================================================
# EXPORT CSV
# =====================================================================
def export_csv(df):
    """
    Export a DataFrame as CSV and allow Streamlit download.
    """
    csv = df.to_csv(index=False).encode('utf-8')
    st.download_button(
        label="Download CSV File",
        data=csv,
        file_name="export_results.csv",
        mime="text/csv"
    )


# =====================================================================
# EXPORT EXCEL
# =====================================================================
def export_excel(df):
    """
    Export a DataFrame as Excel file using BytesIO.
    """
    buffer = BytesIO()
    with pd.ExcelWriter(buffer, engine="xlsxwriter") as writer:
        df.to_excel(writer, index=False, sheet_name="Results")

    st.download_button(
        label="Download Excel File",
        data=buffer,
        file_name="export_results.xlsx",
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )


# =====================================================================
# EXPORT PDF
# =====================================================================
def export_pdf(df):
    """
    Export DataFrame to PDF using reportlab.
    Produces a simple engineering report.
    """

    buffer = BytesIO()
    c = canvas.Canvas(buffer, pagesize=letter)
    width, height = letter

    text = c.beginText(40, height - 40)
    text.setFont("Helvetica", 10)
    text.textLine("BESS Interconnection Engineering Report")
    text.textLine("-----------------------------------------------------")
    text.moveCursor(0, 10)

    for col in df.columns:
        text.textLine(f"{col} : {df[col].iloc[0]}")
        text.moveCursor(0, 10)

    c.drawText(text)
    c.showPage()
    c.save()

    st.download_button(
        label="Download PDF File",
        data=buffer,
        file_name="export_results.pdf",
        mime="application/pdf"
    )


# =====================================================================
# EXPORT DOCX (AUTO-FILL TEMPLATE)
# =====================================================================
def export_docx(summary_dict, template_path="assets/form_templates/template.docx"):
    """
    Fills a DOCX template with key-value pairs from a summary dictionary.
    Dictionary keys should match placeholders in the template.
    Example placeholder: {{POI_MW}}
    """

    try:
        doc = Document(template_path)
    except Exception as e:
        st.error(f"Error loading template: {e}")
        return

    # Replace placeholders
    for p in doc.paragraphs:
        for key, value in summary_dict.items():
            placeholder = "{{" + key + "}}"
            if placeholder in p.text:
                p.text = p.text.replace(placeholder, str(value))

    buffer = BytesIO()
    doc.save(buffer)

    st.download_button(
        label="Download Filled DOCX",
        data=buffer,
        file_name="interconnection_form_filled.docx",
        mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    )
